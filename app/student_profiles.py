"""NCET BCA flexible enrollment and consent-based public student showcase.

Public profiles contain only explicitly published fields. Private email/mobile and
confidential HOD requests are never returned by anonymous API or HTML routes.
"""
from __future__ import annotations

import io
import json
import re
import secrets
import zipfile
from PIL import Image, ImageOps, UnidentifiedImageError
from datetime import datetime
from html import escape
from urllib.parse import urlsplit

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request
from fastapi.responses import HTMLResponse, Response
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, DateTime, ForeignKey, String, Text, LargeBinary, or_, func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session as DbSession, Mapped, mapped_column, defer

from .database import Base, User, AuditLog, utcnow, get_db
from .student_models import BCAStudent
from .student_services import (USN_RE, can_manage_student_roster, student_roster_department,
                               department_scope_values)
from .security import current, csrf_required, hash_password
from .departments import branch_option, branch_name, AcademicDepartment

router = APIRouter(tags=['NCET student enrollment and showcase'])
URL_ROOT = 'https://ncet-verifyflow.onrender.com'
MAX_BATCH = 100
USN_DEPARTMENT_PREFIXES = {'CS': 'Computer Science & Engineering'}

class StudentShowcase(Base):
    __tablename__ = 'bca_student_showcases'
    id: Mapped[int] = mapped_column(primary_key=True)
    student_id: Mapped[int] = mapped_column(ForeignKey('users.id',ondelete='CASCADE'),unique=True,index=True)
    mobile_private: Mapped[str] = mapped_column(String(20),default='')
    email_private: Mapped[str] = mapped_column(String(180),default='')
    about: Mapped[str] = mapped_column(Text,default='')
    skills_json: Mapped[str] = mapped_column(Text,default='[]')
    certificates_json: Mapped[str] = mapped_column(Text,default='[]')
    social_links_json: Mapped[str] = mapped_column(Text,default='{}')
    projects_json: Mapped[str] = mapped_column(Text,default='[]')
    interests_json: Mapped[str] = mapped_column(Text,default='[]')
    published: Mapped[bool] = mapped_column(Boolean,default=False,index=True)
    photo_bytes: Mapped[bytes | None] = mapped_column(LargeBinary,nullable=True)
    photo_mime: Mapped[str | None] = mapped_column(String(32),nullable=True)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True),default=utcnow)

class EnrollmentItem(BaseModel):
    usn: str = Field(min_length=8,max_length=24)
    full_name: str = Field(min_length=3,max_length=130)
    branch: str = Field(min_length=2,max_length=160)
    semester: int = Field(ge=1,le=8)

class EnrollmentBody(BaseModel):
    students: list[EnrollmentItem] = Field(min_length=1,max_length=MAX_BATCH)

class FacultyUpdate(BaseModel):
    full_name: str = Field(min_length=3,max_length=130)
    branch: str = Field(min_length=2,max_length=160)
    semester: int = Field(ge=1,le=8)

class Certificate(BaseModel):
    title: str = Field(min_length=3,max_length=120)
    issuer: str = Field(min_length=2,max_length=120)
    proof_url: str = Field(min_length=12,max_length=650)

class StudentProject(BaseModel):
    title: str = Field(min_length=3,max_length=105)
    description: str = Field(default='',max_length=550)
    technologies: list[str] = Field(default_factory=list,max_length=12)
    github_url: str = Field(default='',max_length=400)
    live_url: str = Field(default='',max_length=400)

class ShowcaseEdit(BaseModel):
    full_name: str = Field(min_length=3,max_length=130)
    mobile: str = Field(default='',max_length=20)
    email: str = Field(default='',max_length=180)
    about: str = Field(default='',max_length=900)
    skills: list[str] = Field(default_factory=list,max_length=100)
    certificates: list[Certificate] = Field(default_factory=list,max_length=8)
    interests: list[str] = Field(default_factory=list,max_length=20)
    projects: list[StudentProject] = Field(default_factory=list,max_length=8)
    social_links: dict[str,str] = Field(default_factory=dict,max_length=12)
    publish: bool = False


SOCIAL_HOSTS={
 'github':'github.com','linkedin':'linkedin.com','leetcode':'leetcode.com',
 'codechef':'codechef.com','hackerrank':'hackerrank.com','kaggle':'kaggle.com',
 'codeforces':'codeforces.com','stackoverflow':'stackoverflow.com',
 'gitlab':'gitlab.com','devto':'dev.to'
}
def validate_social_links(links):
    selected={}
    for platform,raw in links.items():
        if platform not in (*SOCIAL_HOSTS,'portfolio'):
            raise HTTPException(422,'Unsupported coding profile platform.')
        link=raw.strip()
        if not link:continue
        if len(link)>400:raise HTTPException(422,'Profile link is too long.')
        url=urlsplit(link)
        host=(url.hostname or '').lower()
        if url.scheme!='https' or not host or url.username or url.password:
            raise HTTPException(422,'Profile links must use HTTPS.')
        if platform!='portfolio' and host not in (SOCIAL_HOSTS[platform],'www.'+SOCIAL_HOSTS[platform]):
            raise HTTPException(422,'Use the official '+platform+' site for that link.')
        if platform=='portfolio' and (host in ('localhost','127.0.0.1') or host.endswith(('.local','.internal'))):
            raise HTTPException(422,'Portfolio must be publicly accessible.')
        check_public_text(url.path)
        selected[platform]=link
    return selected

def valid_name(name:str)->bool:
    return 3<=len(name)<=130 and name[0] not in '=+-@' and not any(ord(c)<32 for c in name)


def validate_usn_department(usn:str,department:str)->None:
    """Reject the confirmed CS USN code when it is assigned to another department."""
    expected=USN_DEPARTMENT_PREFIXES.get(usn[5:7].upper())
    if expected and department!=expected:
        raise HTTPException(422,f'USN {usn} identifies {expected}; select that department.')


def validate_xlsx_archive(content:bytes)->None:
    """Reject malformed, oversized, or macro-enabled XLSX containers early."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries=archive.infolist()
            if not entries or len(entries)>256:
                raise ValueError('Spreadsheet archive has an invalid number of files.')
            total_size=0
            names=set()
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise ValueError('Encrypted spreadsheets are not supported.')
                if entry.file_size<0 or entry.file_size>4*1024*1024:
                    raise ValueError('Spreadsheet contains an oversized file.')
                total_size+=entry.file_size
                names.add(entry.filename)
            if total_size>8*1024*1024:
                raise ValueError('Spreadsheet expands beyond the 8 MB safety limit.')
            required={'[Content_Types].xml','_rels/.rels','xl/workbook.xml'}
            if not required.issubset(names) or not any(
                name.startswith('xl/worksheets/') and name.endswith('.xml') for name in names):
                raise ValueError('File is not a valid XLSX workbook.')
            if any(name.lower()=='xl/vbaproject.bin' for name in names):
                raise ValueError('Macro-enabled spreadsheets are not supported.')
    except (zipfile.BadZipFile,zipfile.LargeZipFile) as exc:
        raise ValueError('Spreadsheet archive is damaged.') from exc


def validate_enrollment(db:DbSession,actor:User,items:list[EnrollmentItem])->list[tuple[str,str,str,int]]:
    result=[];seen=set()
    for i,item in enumerate(items,1):
        usn=item.usn.strip().upper()
        name=item.full_name.strip()
        if not USN_RE.fullmatch(usn):raise HTTPException(422,f'Invalid USN at entry {i}; expected 1NC followed by letters or digits.')
        if not valid_name(name):raise HTTPException(422,f'Student name is required at entry {i}.')
        if usn in seen:raise HTTPException(422,f'Duplicate USN at entry {i}.')
        branch=branch_option(db,item.branch).code
        validate_usn_department(usn,branch)
        if actor.role!='admin' and branch!=student_roster_department(db,actor):
            raise HTTPException(403,'Department faculty may register students only in their assigned branch.')
        seen.add(usn);result.append((usn,name,branch,item.semester))
    return result


def create_students(db, actor, items:list[EnrollmentItem]):
    can_manage_student_roster(actor)
    values=validate_enrollment(db,actor,items)
    if db.query(User.id).filter(User.username.in_([u.lower() for u,n,b,s in values])).first() or \
       db.query(BCAStudent.id).filter(BCAStudent.usn.in_([u for u,n,b,s in values])).first():
        raise HTTPException(409,'One or more USNs are already enrolled. No accounts were created.')
    creds=[]
    try:
        for usn,name,branch,semester in values:
            password=secrets.token_urlsafe(24)
            # This is deliberately not a guessable USN-derived password.
            user=User(username=usn.lower(),display_name=name,
                      role='student',department=branch,
                      password_hash=hash_password(password),active=True)
            db.add(user);db.flush()
            db.add(BCAStudent(user_id=user.id,usn=usn,semester=semester,must_change_password=True))
            creds.append({'usn':usn,'name':name,'branch':branch_name(db,branch),'semester':semester,'temporary_password':password})
        db.add(AuditLog(actor_id=actor.id,event='bca.flexible_enrollment',detail=f'{len(values)} new USNs'))
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409,'Duplicate student account: the batch was not imported.')
    return {'created':len(creds),'credentials':creds}


@router.post('/api/bca/students/enroll')
def enroll(body:EnrollmentBody,actor:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    return create_students(db,actor,body.students)


@router.post('/api/bca/students/import-excel')
async def import_excel(file:UploadFile=File(...),actor:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    can_manage_student_roster(actor)
    if not (file.filename or '').lower().endswith('.xlsx'):
        raise HTTPException(422,'Upload an .xlsx spreadsheet.')
    content=await file.read(350_001)
    if len(content)>350_000:raise HTTPException(413,'Spreadsheet must be smaller than 350 KB.')
    try:
        validate_xlsx_archive(content)
    except ValueError as exc:
        raise HTTPException(422,'Invalid Excel roster: '+str(exc)) from exc
    from openpyxl.xml import DEFUSEDXML
    if not DEFUSEDXML:
        raise HTTPException(503,'Safe Excel XML parsing is unavailable; contact the administrator.')
    from openpyxl import load_workbook
    try:
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        sheet=book.active
        if ((sheet.max_row is not None and sheet.max_row>101) or
                (sheet.max_column is not None and sheet.max_column>4)):
            raise ValueError('Use exactly four columns and at most 100 student rows.')
        iterator=sheet.iter_rows(min_row=1,max_row=102,max_col=5,values_only=True)
        header_row=next(iterator)
        if len(header_row)>4 and header_row[4] is not None:
            raise ValueError('Use exactly four columns: USN, Name, Branch, Sem.')
        header=tuple(str(c or '').strip().lower() for c in header_row[:4])
        aliases={'name':'full_name','department':'branch','sem':'semester'}
        normalized=tuple(aliases.get(h,h) for h in header)
        if normalized!=('usn','full_name','branch','semester'):
            raise ValueError('Columns must be USN, Name, Branch, Sem in that order.')
        parsed=[]
        for row in iterator:
            if not row or not any(c is not None and str(c).strip() for c in row):continue
            if len(row)>4 and row[4] is not None:
                raise ValueError('Use exactly four columns: USN, Name, Branch, Sem.')
            if len(parsed)>=100:raise ValueError('Maximum 100 student rows per spreadsheet.')
            if any(c is None or not str(c).strip() for c in row[:4]):
                raise ValueError('All four fields are required for each student.')
            try:semester=int(row[3])
            except (TypeError,ValueError):raise ValueError('Semester must be a whole number between 1 and 8.')
            parsed.append(EnrollmentItem(usn=str(row[0]).strip(),
                                         full_name=str(row[1]).strip(),
                                         branch=str(row[2]).strip(),semester=semester))
    except Exception as exc:
        raise HTTPException(422,'Invalid Excel roster. Use required USN, Name, Branch, Sem (1–8).') from exc
    finally:
        if 'book' in locals():book.close()
    if not parsed:raise HTTPException(422,'The spreadsheet is empty.')
    return create_students(db,actor,parsed)


@router.patch('/api/bca/students/{usn}')
def faculty_edit(usn:str,body:FacultyUpdate,actor:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    can_manage_student_roster(actor)
    found=(db.query(User,BCAStudent).join(BCAStudent,BCAStudent.user_id==User.id)
           .filter(BCAStudent.usn==usn.strip().upper()).first())
    if not found:raise HTTPException(404,'Student USN not found.')
    name=body.full_name.strip()
    if not valid_name(name):raise HTTPException(422,'Enter a valid student name.')
    user,student=found
    target=branch_option(db,body.branch)
    validate_usn_department(student.usn,target.code)
    scope=department_scope_values(db,actor.department)
    if actor.role!='admin' and (user.department not in scope or target.code!=student_roster_department(db,actor)):
        raise HTTPException(403,'Faculty cannot transfer students between departments.')
    user.display_name=name;student.semester=body.semester;user.department=target.code
    # Changing official enrollment data requires the student to explicitly republish.
    showcase=db.query(StudentShowcase).filter(StudentShowcase.student_id==user.id).first()
    if showcase:showcase.published=False
    db.add(AuditLog(actor_id=actor.id,event='bca.student_details_updated',detail=student.usn))
    db.commit()
    return {'ok':True,'branch':user.department,'publication_review_required':bool(showcase)}


def _my_student(db,user):
    if user.role!='student':raise HTTPException(403,'Student account required.')
    found=db.query(BCAStudent).filter(BCAStudent.user_id==user.id).first()
    if not found:raise HTTPException(403,'Enrollment record missing.')
    return found


PUBLIC_EMAIL_PATTERN=re.compile(r'[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}')
PUBLIC_PHONE_PATTERN=re.compile(r'(?<!\w)\+?\d[\d() .\-]{7,}\d(?!\w)')


def check_public_text(text:str)->None:
    if PUBLIC_EMAIL_PATTERN.search(text) or PUBLIC_PHONE_PATTERN.search(text):
        raise HTTPException(422,'Remove email addresses and phone numbers from fields published publicly.')


def validate_proof(url:str)->str:
    url=url.strip()
    parts=urlsplit(url)
    if parts.scheme!='https' or not parts.hostname or parts.username or parts.password:
        raise HTTPException(422,'Certificate proof must be a public HTTPS URL.')
    if parts.hostname.lower() in ('localhost','127.0.0.1','0.0.0.0') or parts.hostname.endswith('.local'):  # nosec B104 - reject list, not socket bind
        raise HTTPException(422,'Invalid certificate host.')
    return url


def validate_public_link(raw:str,*,github_only:bool=False)->str:
    url=raw.strip()
    parsed=urlsplit(url)
    host=(parsed.hostname or '').lower()
    if (parsed.scheme!='https' or not host or parsed.username or parsed.password or
        host in ('localhost','127.0.0.1','0.0.0.0','::1') or  # nosec B104 - reject list, not socket bind
        host.endswith(('.local','.internal','.localhost')) or
        any(ord(c)<32 for c in url)):
        raise HTTPException(422,'Project links must be valid public HTTPS URLs.')
    if github_only and host not in ('github.com','www.github.com'):
        raise HTTPException(422,'Project GitHub URL must point to github.com.')
    check_public_text(parsed.path)
    return url


def _showcase(db,user):
    return db.query(StudentShowcase).filter(StudentShowcase.student_id==user.id).first()


@router.get('/api/bca/showcase/me')
def my_showcase(user:User=Depends(current),db:DbSession=Depends(get_db)):
    student=_my_student(db,user)
    s=_showcase(db,user)
    return {'usn':student.usn,'full_name':user.display_name if user.display_name!='Registration pending' else '',
            'department':branch_name(db,user.department),'semester':student.semester,
            'must_change_password':student.must_change_password,
            'mobile':s.mobile_private if s else '', 'email':s.email_private if s else '',
            'about':s.about if s else '', 'skills':json.loads(s.skills_json) if s else [],
            'certificates':json.loads(s.certificates_json) if s else [],
            'projects':json.loads(s.projects_json or '[]') if s else [],
            'interests':json.loads(s.interests_json or '[]') if s else [],
            'social_links':json.loads(s.social_links_json or '{}') if s else {},
            'published':s.published if s else False,
            'photo_available':bool(s.photo_bytes) if s else False}


@router.put('/api/bca/showcase/me')
def edit_showcase(body:ShowcaseEdit,user:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    student=_my_student(db,user)
    if student.must_change_password:raise HTTPException(403,'Change your temporary password before editing your profile.')
    name=body.full_name.strip()
    if not valid_name(name):raise HTTPException(422,'Enter your complete name.')
    phone=body.mobile.strip()
    if phone and not re.fullmatch(r'\+?[0-9 ]{8,20}',phone):
        raise HTTPException(422,'Enter a valid private mobile number.')
    email=body.email.strip().lower()
    if email and (not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+',email) or len(email)>180):
        raise HTTPException(422,'Enter a valid private email address.')
    skills=[]
    for item in body.skills:
        value=item.strip()
        if not 1<=len(value)<=80 or any(ord(c)<32 for c in value):
            raise HTTPException(422,'Each skill must contain 1–80 printable characters.')
        if value.casefold() not in [x.casefold() for x in skills]:skills.append(value)
    if sum(len(v) for v in skills)>5000:raise HTTPException(422,'All skills combined must be under 5,000 characters.')
    social_links=validate_social_links(body.social_links)
    interests=[]
    for item in body.interests:
        interest=item.strip()
        if not 2<=len(interest)<=75 or any(ord(c)<32 for c in interest):
            raise HTTPException(422,'Each interest must have 2–75 printable characters.')
        check_public_text(interest)
        if interest.casefold() not in [v.casefold() for v in interests]:interests.append(interest)
    projects=[]
    for item in body.projects:
        title=item.title.strip();description=item.description.strip()
        if not 3<=len(title)<=105:raise HTTPException(422,'Project titles must contain 3–105 characters.')
        check_public_text(title);check_public_text(description)
        tech=[]
        for value in item.technologies:
            entry=value.strip()
            if not 1<=len(entry)<=50:raise HTTPException(422,'Project technologies must be 1–50 characters.')
            check_public_text(entry)
            if entry.casefold() not in [v.casefold() for v in tech]:tech.append(entry)
        github=validate_public_link(item.github_url,github_only=True) if item.github_url else ''
        live=validate_public_link(item.live_url) if item.live_url else ''
        projects.append({'title':title,'description':description,'technologies':tech,
                         'github_url':github,'live_url':live})
    check_public_text(name)
    check_public_text(body.about)
    for value in skills:check_public_text(value)
    certificates=[]
    for certificate in body.certificates:
        check_public_text(certificate.title)
        check_public_text(certificate.issuer)
        check_public_text(certificate.proof_url.split('?',1)[-1] if '?' in certificate.proof_url else '')
        certificates.append({'title':certificate.title.strip(),'issuer':certificate.issuer.strip(),
                            'proof_url':validate_proof(certificate.proof_url)})
    profile=_showcase(db,user)
    if not profile:
        profile=StudentShowcase(student_id=user.id)
        db.add(profile)
    user.display_name=name
    profile.mobile_private=phone;profile.email_private=email
    profile.about=body.about.strip()
    profile.skills_json=json.dumps(skills)
    profile.certificates_json=json.dumps(certificates)
    profile.social_links_json=json.dumps(social_links)
    profile.projects_json=json.dumps(projects)
    profile.interests_json=json.dumps(interests)
    profile.published=bool(body.publish)
    profile.updated_at=utcnow()
    db.add(AuditLog(actor_id=user.id,event='bca.showcase_updated',
                    detail=f'{student.usn}: {"public" if profile.published else "private"}'))
    db.commit()
    return {'ok':True,'published':profile.published,'public_url':f'/students/profile/{student.usn}' if profile.published else None}


PHOTO_LIMIT=150_000  # uploaded file must be strictly smaller than 150 KB (decimal)
ALLOWED_PHOTO_FORMATS={'PNG','JPEG','WEBP'}


@router.put('/api/bca/showcase/me/photo')
async def upload_profile_photo(
    file:UploadFile=File(...),
    user:User=Depends(csrf_required),
    db:DbSession=Depends(get_db),
):
    student=_my_student(db,user)
    if student.must_change_password:
        raise HTTPException(403,'Change your temporary password before uploading a photo.')
    if file.content_type not in ('image/png','image/jpeg','image/webp'):
        raise HTTPException(422,'Choose a PNG, JPEG or WebP photo under 150 KB.')
    raw=await file.read(PHOTO_LIMIT+1)
    await file.close()
    if not raw or len(raw)>=PHOTO_LIMIT:
        raise HTTPException(413,'Profile photo must be smaller than 150 KB.')
    try:
        with Image.open(io.BytesIO(raw)) as incoming:
            if incoming.format not in ALLOWED_PHOTO_FORMATS:
                raise HTTPException(422,'Unsupported photo format.')
            if incoming.width>2500 or incoming.height>2500 or incoming.width<32 or incoming.height<32:
                raise HTTPException(422,'Photo dimensions must be between 32 and 2500 pixels.')
            incoming.load()
            image=ImageOps.exif_transpose(incoming).convert('RGB')
            image.thumbnail((640,640),Image.Resampling.LANCZOS)
            photo=None
            for quality in (83,72,60,48):
                out=io.BytesIO()
                image.save(out,format='WEBP',quality=quality,method=4)
                if len(out.getvalue())<PHOTO_LIMIT:
                    photo=out.getvalue()
                    break
            if photo is None:
                raise HTTPException(422,'Photo could not fit under 150 KB. Choose a smaller image.')
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError):
        raise HTTPException(422,'Invalid or damaged image.')
    profile=_showcase(db,user)
    if not profile:
        profile=StudentShowcase(student_id=user.id)
        db.add(profile)
    profile.photo_bytes=photo
    profile.photo_mime='image/webp'
    profile.updated_at=utcnow()
    db.add(AuditLog(actor_id=user.id,event='bca.showcase_photo_updated',detail=student.usn))
    db.commit()
    return {'ok':True,'photo_available':True,'size_bytes':len(photo),
            'public_photo_url':f'/students/profile/{student.usn}/photo' if profile.published else None}


@router.get('/api/bca/showcase/me/photo')
def private_profile_photo(user:User=Depends(current),db:DbSession=Depends(get_db)):
    _my_student(db,user)
    profile=_showcase(db,user)
    if not profile or not profile.photo_bytes:
        raise HTTPException(404,'Profile photo not uploaded.')
    return Response(profile.photo_bytes,media_type='image/webp',
                    headers={'Cache-Control':'private,no-store','X-Robots-Tag':'noindex,nofollow'})


@router.delete('/api/bca/showcase/me/photo')
def delete_profile_photo(user:User=Depends(csrf_required),db:DbSession=Depends(get_db)):
    student=_my_student(db,user)
    profile=_showcase(db,user)
    if not profile or not profile.photo_bytes:
        raise HTTPException(404,'Profile photo not uploaded.')
    profile.photo_bytes=None
    profile.photo_mime=None
    profile.updated_at=utcnow()
    db.add(AuditLog(actor_id=user.id,event='bca.showcase_photo_removed',detail=student.usn))
    db.commit()
    return {'ok':True,'photo_available':False}


@router.get('/students/profile/{usn}/photo',include_in_schema=False)
def published_profile_photo(usn:str,db:DbSession=Depends(get_db)):
    if not USN_RE.fullmatch(usn.strip().upper()):
        raise HTTPException(404,'Profile unavailable.')
    row=public_query(db).filter(BCAStudent.usn==usn.strip().upper()).first()
    if not row or not row[2].photo_bytes:
        raise HTTPException(404,'Profile photo unavailable.')
    return Response(row[2].photo_bytes,media_type='image/webp',
                    headers={'Cache-Control':'no-store',
                             'X-Content-Type-Options':'nosniff'})


def public_query(db):
    return (db.query(User,BCAStudent,StudentShowcase)
            .join(BCAStudent,BCAStudent.user_id==User.id)
            .join(StudentShowcase,StudentShowcase.student_id==User.id)
            .filter(User.active.is_(True),BCAStudent.must_change_password.is_(False),
                    StudentShowcase.published.is_(True)))


def public_record(user,student,showcase):
    """Deliberate allowlist; private mobile and email never leave this function."""
    return {'name':user.display_name,'usn':student.usn,
            'department':user.department,'semester':student.semester,
            'about':showcase.about,'skills':json.loads(showcase.skills_json),
            'certificates':json.loads(showcase.certificates_json),
            'projects':json.loads(showcase.projects_json or '[]'),
            'interests':json.loads(showcase.interests_json or '[]'),
            'social_links':json.loads(showcase.social_links_json or '{}'),
            'photo_url':f'/students/profile/{student.usn}/photo' if showcase.photo_mime else None,
            'updated_at':showcase.updated_at.isoformat() if showcase.updated_at else None}


@router.get('/api/bca/directory')
def public_directory(q:str='',limit:int=30,offset:int=0,db:DbSession=Depends(get_db)):
    if not 1<=limit<=50 or not 0<=offset<=1_000_000 or len(q)>60:raise HTTPException(422,'Invalid search.')
    query=public_query(db)
    if q.strip():
        term='%'+q.strip().replace('%','\\%').replace('_','\\_')+'%'
        query=query.filter(or_(BCAStudent.usn.ilike(term,escape='\\'),
                                User.display_name.ilike(term,escape='\\')))
    total=query.with_entities(func.count(StudentShowcase.id)).scalar() or 0
    # Never download up to 50 stored images just to render a public roster.
    rows=(query.options(defer(StudentShowcase.photo_bytes))
          .order_by(BCAStudent.usn).offset(offset).limit(limit).all())
    codes={u.department for u,_,_ in rows}
    names=dict(db.query(AcademicDepartment.code,AcademicDepartment.name)
               .filter(AcademicDepartment.code.in_(codes)).all()) if codes else {}
    return {'total':total,'offset':offset,'limit':limit,'items':[{'name':u.display_name,'usn':p.usn,'department':names.get(u.department,u.department),
                       'skills':json.loads(s.skills_json),
                       'photo_url':f'/students/profile/{p.usn}/photo' if s.photo_mime else None} for u,p,s in rows]}


@router.get('/api/bca/directory/{usn}')
def public_one(usn:str,db:DbSession=Depends(get_db)):
    if not USN_RE.fullmatch(usn.strip().upper()):raise HTTPException(404,'Profile unavailable.')
    record=(public_query(db).options(defer(StudentShowcase.photo_bytes))
            .filter(BCAStudent.usn==usn.strip().upper()).first())
    if not record:raise HTTPException(404,'Profile unavailable.')
    result=public_record(*record)
    result['department']=branch_name(db,record[0].department)
    return result


@router.get('/students/profile/{usn}',response_class=HTMLResponse,include_in_schema=False)
def public_profile_html(usn:str,db:DbSession=Depends(get_db)):
    """Institutional, résumé-style single-column portfolio; safe for search engines."""
    info=public_one(usn,db)
    name=escape(info['name'])
    sid=escape(info['usn'])
    branch=escape(info['department'])
    summary=escape((info['about'] or f"{info['name']} studies at Nagarjuna College of Engineering and Technology.")[:155],quote=True)
    canonical=f"{URL_ROOT}/students/profile/{sid}"
    photo=info.get('photo_url')
    absolute_photo=URL_ROOT+photo if photo else ''
    initials=escape(''.join(part[0] for part in info['name'].split()[:2]).upper() or 'NCET')
    avatar=(f'<img class="resume-avatar" src="{escape(photo,quote=True)}" alt="{name} profile photograph" width="108" height="108">'
            if photo else f'<div class="resume-initials" aria-label="Profile photo not provided">{initials}</div>')
    def chips(values):
        return ''.join(f'<li>{escape(x)}</li>' for x in values)
    def empty(text):
        return '<p class="resume-empty">'+escape(text)+'</p>'
    skill_section='<ul class="resume-chips">'+chips(info['skills'])+'</ul>' if info['skills'] else empty('Skills will appear when the student adds them.')
    interests='<ul class="resume-chips resume-interests">'+chips(info['interests'])+'</ul>' if info['interests'] else empty('Academic interests have not been shared yet.')
    def link(url,title,klass='resume-text-link'):
        return (f'<a class="{klass}" href="{escape(url,quote=True)}" target="_blank" rel="noopener noreferrer nofollow">'
                f'{escape(title)} <span aria-hidden="true">↗</span></a>')
    projects=[]
    for project in info['projects']:
        links=''.join(filter(None,[
            link(project.get('github_url',''),'GitHub source') if project.get('github_url') else '',
            link(project.get('live_url',''),'Live project') if project.get('live_url') else ''
        ]))
        tags=('<ul class="resume-chips resume-project-tags">'+chips(project.get('technologies',[]))+'</ul>') if project.get('technologies') else ''
        projects.append('<article class="resume-project"><h3>'+escape(project['title'])+'</h3>'+
                        '<p>'+escape(project.get('description') or 'Student project')+'</p>'+tags+
                        ('<div class="resume-links">'+links+'</div>' if links else '')+'</article>')
    project_html=''.join(projects) if projects else empty('Student projects have not been added yet.')
    certificates=[]
    for cert in info['certificates']:
        certificates.append('<li class="resume-certificate"><div><strong>'+escape(cert['title'])+
                            '</strong><p>'+escape(cert['issuer'])+'</p></div>'+
                            link(cert['proof_url'],'View credential')+'</li>')
    cert_html='<ul class="resume-certificates">'+''.join(certificates)+'</ul>' if certificates else empty('No certificates shared yet.')
    socials=''.join('<li>'+link(url,platform.title(),'resume-social')+'</li>' for platform,url in info['social_links'].items())
    social_html='<ul class="resume-socials">'+socials+'</ul>' if socials else empty('Public coding and professional links have not been added yet.')
    projects_schema=[{'@type':'CreativeWork','name':p['title'],
                      'description':p.get('description',''),
                      'url':p.get('live_url') or p.get('github_url') or canonical}
                     for p in info['projects']]
    schema={
        '@context':'https://schema.org','@type':'ProfilePage','url':canonical,
        'dateModified':info.get('updated_at'),
        'about':{'@type':'Person','name':info['name'],'identifier':info['usn'],
                 'description':info['about'] or f"{info['name']} studies at NCET.",
                 'affiliation':{'@type':'CollegeOrUniversity',
                                'name':'Nagarjuna College of Engineering and Technology'},
                 'knowsAbout':info['skills']+info['interests'],
                 'sameAs':list(info['social_links'].values())},
        'mainEntity':{'@type':'Person','name':info['name'],'identifier':info['usn']},
        'hasPart':projects_schema,
    }
    if absolute_photo:schema['about']['image']=absolute_photo
    json_ld=json.dumps(schema,ensure_ascii=False).replace('<','\\u003c')
    og_image=(f'<meta property="og:image" content="{escape(absolute_photo,quote=True)}">'
              if absolute_photo else '<meta property="og:image" content="'+URL_ROOT+'/static/ncet-brand.png">')
    return HTMLResponse(f'''<!doctype html><html lang="en"><head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="theme-color" content="#0E2455">
<title>{name} ({sid}) | NCET Student Portfolio</title>
<meta name="description" content="{summary}">
<meta name="robots" content="index,follow,max-image-preview:large">
<meta property="og:type" content="profile"><meta property="og:title" content="{name} — NCET Student Portfolio">
<meta property="og:description" content="{summary}"><meta property="og:url" content="{canonical}">
{og_image}<link rel="canonical" href="{canonical}">
<link rel="icon" href="/static/ncet-brand.svg">
<script type="application/ld+json">{json_ld}</script>
<link rel="preload" as="image" href="/static/ncet-brand.svg" fetchpriority="high">
<link rel="stylesheet" href="/static/resume-profile.css?v=20261017-2">
<link rel="stylesheet" href="/static/institutional-ui.css?v=20261017d">
<script src="/static/student-profile-tabs.js?v=20261017" defer></script>
</head><body class="ncet-profile-page">
<a class="skip-link" href="#student-resume">Skip to portfolio</a>
<header class="resume-site-header">
<div class="resume-site-inner">
<a class="resume-college" href="/students" aria-label="NCET Student Portal">
<img src="/static/ncet-brand.svg" alt="Nagarjuna College of Engineering and Technology" width="178" height="88" fetchpriority="high" decoding="async">
<span>Student portfolios<small>Official NCET student portal</small></span></a>
<nav aria-label="Profile navigation"><a href="/students/directory">Student directory</a><a href="/students">Student Portal</a></nav>
<img class="resume-group-logo" src="/static/ngi-brand.svg" alt="Nagarjuna Group of Institutions" width="78" height="70" decoding="async">
</div></header>
<main id="student-resume" class="resume-container">
<div class="resume-paper">
<div class="resume-accreditation"><span>NCET · STUDENT PORTFOLIO</span>
<img src="/static/naac-brand.svg" alt="Institutional accreditation: NAAC A+" width="61" height="61" decoding="async"></div>
<div class="resume-heading">{avatar}<div class="resume-heading-copy">
<p class="resume-overline">Student profile · Voluntarily published</p>
<h1>{name}</h1><p class="resume-subtitle">{branch} <span aria-hidden="true">·</span> Semester {info['semester']}</p>
<p class="resume-usn"><strong>USN</strong> {sid}</p></div></div>
<nav class="resume-anchors" aria-label="Portfolio sections">
<a href="#about" id="profile-tab-about">About</a><a href="#skills" id="profile-tab-skills">Skills</a><a href="#projects" id="profile-tab-projects">Projects</a>
<a href="#certifications" id="profile-tab-certifications">Certifications</a><a href="#interests" id="profile-tab-interests">Interests</a><a href="#links" id="profile-tab-links">Links</a>
</nav>
<div class="resume-tab-stage" id="resume-tab-stage" aria-live="off">
<section id="about" class="resume-section"><h2>About</h2>
<p class="resume-about">{escape(info['about']) or 'This student has not yet added a summary.'}</p></section>
<section id="skills" class="resume-section"><h2>Technical skills</h2>{skill_section}</section>
<section id="projects" class="resume-section"><h2>Projects</h2>{project_html}</section>
<section id="certifications" class="resume-section"><h2>Certifications</h2>{cert_html}</section>
<section id="interests" class="resume-section"><h2>Areas of interest</h2>{interests}</section>
<section id="links" class="resume-section"><h2>GitHub, live work & professional links</h2>{social_html}</section>
</div>
<div class="resume-disclaimer"><p>Published by the student. NCET has not independently verified personal projects, skills or linked credentials.</p>
<a href="/students/directory">← Browse student directory</a></div>
</div>
</main>
<footer class="resume-footer"><img src="/static/ngi-brand.svg" alt="Nagarjuna Group of Institutions" width="53" height="47" loading="lazy" decoding="async">
<span>NCET VerifyFlow · Academic collaboration & student success</span>
<a href="/students">Student Portal</a></footer>
<script src="/static/institutional-ui.js?v=20261017d" defer></script>
</body></html>''',headers={'Cache-Control':'no-store'})


@router.get('/students/directory',include_in_schema=False)
def directory_page(page:int=1,db:DbSession=Depends(get_db)):
    """Small paginated server-rendered lists, with every published profile crawlable."""
    from pathlib import Path
    if page<1 or page>100000:raise HTTPException(404,'Directory page not found.')
    static=Path(__file__).resolve().parents[1]/'static'
    html=(static/'student-directory.html').read_text(encoding='utf-8')
    per_page=18
    total=public_query(db).with_entities(func.count(StudentShowcase.id)).scalar() or 0
    pages=max(1,(total+per_page-1)//per_page)
    if page>pages:raise HTTPException(404,'Directory page not found.')
    rows=(public_query(db).options(defer(StudentShowcase.photo_bytes))
          .order_by(BCAStudent.usn).offset((page-1)*per_page).limit(per_page).all())
    items=''.join(
        '<li><a href="/students/profile/'+escape(student.usn,quote=True)+'">'+
        escape(user.display_name)+'</a> <span class="muted">'+escape(student.usn)+'</span></li>'
        for user,student,showcase in rows
    )
    controls=(
        ('<a class="btn" href="/students/directory?page='+str(page-1)+'">← Previous</a>' if page>1 else '')+
        '<span>Page '+str(page)+' of '+str(pages)+'</span>'+
        ('<a class="btn" href="/students/directory?page='+str(page+1)+'">Next →</a>' if page<pages else '')
    )
    section=('<section class="panel server-profile-links" aria-label="Published student profiles">'+
             '<h2>Browse all published portfolios</h2>'+
             '<p class="small muted">'+str(total)+' student-published profiles. Showing '+str(len(rows))+' on this page.</p>'+
             '<ul>'+items+'</ul><nav class="profile-pagination" aria-label="Directory pages">'+controls+'</nav></section>')
    html=html.replace('</main>',section+'</main>')
    if page>1:
        html=html.replace('href="https://ncet-verifyflow.onrender.com/students/directory"',
                          'href="https://ncet-verifyflow.onrender.com/students/directory?page='+str(page)+'"')
    return HTMLResponse(html,headers={'Cache-Control':'no-store'})



def published_sitemap_paths(db:DbSession)->list[str]:
    return ['/students/directory'] + [f'/students/profile/{p.usn}'
      for p in (db.query(BCAStudent).join(User,User.id==BCAStudent.user_id)
          .join(StudentShowcase,StudentShowcase.student_id==User.id)
          .filter(User.active.is_(True),BCAStudent.must_change_password.is_(False),
                  StudentShowcase.published.is_(True)).order_by(BCAStudent.usn).limit(10000).all())]
